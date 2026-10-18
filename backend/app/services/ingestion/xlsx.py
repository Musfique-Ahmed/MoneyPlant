"""XLSX ingestion. Reads values only — never evaluates formulas."""

from __future__ import annotations

import io
from typing import Any

from app.services.ingestion.csv import ValidationIssue, suggest_mapping


def preview_xlsx(content: bytes, max_rows: int = 50) -> dict[str, Any]:
    """Return a preview dict for an XLSX file."""
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # noqa: BLE001
        return {
            "headers": [],
            "sample_rows": [],
            "suggested_mapping": {},
            "row_count": 0,
            "sheet_name": None,
            "encoding": "binary",
            "issues": [
                ValidationIssue(
                    row=0, field="file", severity="error", message=f"openpyxl not available: {exc}"
                ).to_dict()
            ],
        }

    # ``data_only=True`` reads the cached value of a formula rather than
    # evaluating it. We never let a workbook formula run.
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    headers: list[str] = []
    sample_rows: list[dict[str, str]] = []
    total = 0
    for i, row in enumerate(rows_iter, start=1):
        if i == 1:
            headers = [str(c) if c is not None else "" for c in row]
            continue
        total += 1
        if i - 1 <= max_rows:
            sample_rows.append(
                {headers[j] if j < len(headers) else f"col{j}": ("" if v is None else str(v)) for j, v in enumerate(row)}
            )
    return {
        "headers": headers,
        "sample_rows": sample_rows,
        "suggested_mapping": suggest_mapping(headers),
        "row_count": total,
        "sheet_name": sheet_name,
        "encoding": "binary",
        "issues": [],
    }


__all__ = ["preview_xlsx"]